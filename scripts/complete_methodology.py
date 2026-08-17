from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\Ramos\Downloads\Fase_1.docx")
OUTPUT = Path(r"C:\xampp\htdocs\Sistema-de-Facturaci-n-INS-\output\documents\Fase_1_metodologia_completada.docx")


def insert_paragraph_after(anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    if style in {"List Bullet", "List Number"}:
        paragraph.style = "List Paragraph"
        properties = paragraph._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number_id = OxmlElement("w:numId")
        number_id.set(qn("w:val"), "3" if style == "List Bullet" else "1")
        numbering.append(level)
        numbering.append(number_id)
        properties.append(numbering)
    else:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


document = Document(SOURCE)

for paragraph in document.paragraphs:
    if "Asimismo, se utilizará Bootstrap" in paragraph.text:
        for run in paragraph.runs:
            if "Bootstrap" in run.text:
                run.text = run.text.replace("Bootstrap", "Tailwind CSS")

methodology_heading = next(
    paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Metodología de trabajo."
)

content = [
    ("Metodología para el desarrollo de software", "Heading 2"),
    (
        "Para el desarrollo del sistema se aplicará Scrum de forma adaptada al contexto académico del equipo. Esta metodología ágil permitirá organizar el trabajo en iteraciones cortas, priorizar las funcionalidades de mayor valor y obtener incrementos funcionales que puedan revisarse periódicamente. Su aplicación facilitará la distribución de responsabilidades entre los integrantes, el seguimiento del avance y la incorporación de mejoras a partir de la retroalimentación recibida.",
        "Normal",
    ),
    (
        "El trabajo se dividirá en sprints de dos semanas. Al inicio de cada sprint se seleccionarán las historias de usuario que serán desarrolladas; durante la iteración se dará seguimiento a las tareas y, al finalizar, se presentará el incremento obtenido para su revisión. Las tareas se registrarán en un tablero de gestión con los estados pendiente, en proceso, en revisión y completado.",
        "Normal",
    ),
    ("Actividades principales de Scrum", "Heading 2"),
    ("Planificación del sprint: selección y estimación de las funcionalidades que se desarrollarán durante la iteración.", "List Bullet"),
    ("Reunión de seguimiento: revisión breve del trabajo realizado, las actividades pendientes y los impedimentos encontrados.", "List Bullet"),
    ("Revisión del sprint: demostración del incremento funcional y validación de su correspondencia con los requerimientos.", "List Bullet"),
    ("Retrospectiva: identificación de aspectos positivos, dificultades y acciones de mejora para la siguiente iteración.", "List Bullet"),
    ("Artefactos de trabajo", "Heading 2"),
    ("Product Backlog: lista priorizada de requerimientos, historias de usuario y mejoras del sistema.", "List Bullet"),
    ("Sprint Backlog: conjunto de tareas seleccionadas para ser desarrolladas durante cada sprint.", "List Bullet"),
    ("Incremento: versión funcional del sistema obtenida al finalizar cada iteración.", "List Bullet"),
    ("Propuesta para la creación del prototipo", "Heading 2"),
    (
        "Se empleará un prototipo evolutivo de alta fidelidad. A diferencia de un modelo únicamente visual, este prototipo estará conectado a la base de datos y permitirá ejecutar los procesos principales del sistema. La primera versión representará al menos el 30 % de la aplicación y servirá como base para incorporar progresivamente el resto de las funcionalidades hasta obtener el producto final.",
        "Normal",
    ),
    (
        "El prototipo inicial incluye un panel administrativo, gestión de clientes, categorías y productos, control básico de existencias, alertas de inventario bajo, registro de ventas, cálculo automático de subtotal, descuento, impuesto y total, actualización del inventario, historial de ventas y visualización de facturas imprimibles. En iteraciones posteriores se incorporarán la autenticación, los perfiles de administrador y vendedor, permisos, movimientos detallados de inventario, reportes y pruebas adicionales.",
        "Normal",
    ),
    ("Roles asignados", "Heading 2"),
    ("Product Owner - Karla Lissette Mejía Ortiz: representa las necesidades de los usuarios, prioriza los requerimientos y valida los incrementos del sistema.", "List Bullet"),
    ("Scrum Master - José Alberto Martínez Núñez: facilita la metodología, organiza las reuniones y ayuda a resolver impedimentos del equipo.", "List Bullet"),
    ("Desarrollo y prototipado - William Antonio Ramos Rodríguez: implementa la interfaz, los módulos funcionales y la integración del prototipo en Laravel.", "List Bullet"),
    ("Análisis, documentación y pruebas - Priscila Lisseth Osorio Escobar: apoya el levantamiento de requerimientos, documenta resultados y verifica los criterios de aceptación.", "List Bullet"),
    ("Base de datos y control de versiones - Jefry Daniel Torres Navarro: diseña y revisa la estructura de datos, apoya la integración y supervisa el repositorio de código.", "List Bullet"),
    (
        "Aunque se establecen responsabilidades principales, el equipo de desarrollo trabajará de forma colaborativa. Los integrantes podrán apoyar otras actividades de acuerdo con las necesidades de cada sprint, manteniendo evidencia individual de sus aportes mediante commits y tareas asignadas.",
        "Normal",
    ),
    ("Plan de iteraciones", "Heading 2"),
    ("Sprint 0 - Preparación: análisis de requerimientos, historias de usuario, diagramas, configuración de Laravel, MySQL y repositorio.", "List Number"),
    ("Sprint 1 - Prototipo de Fase 1: panel, clientes, categorías, productos, inventario básico, ventas y factura imprimible.", "List Number"),
    ("Sprint 2 - Seguridad y operaciones: autenticación, roles, permisos y movimientos detallados de inventario.", "List Number"),
    ("Sprint 3 - Seguimiento: reportes, consultas, mejoras de usabilidad y validaciones adicionales.", "List Number"),
    ("Sprint 4 - Cierre: pruebas integrales, corrección de errores, manual de usuario, video y preparación de la entrega final.", "List Number"),
    ("Justificación de la metodología, tecnologías y herramientas", "Heading 2"),
    (
        "Se seleccionó Scrum porque el proyecto puede dividirse en módulos independientes y entregarse mediante incrementos verificables. Esta metodología permite responder a cambios en los requerimientos, hacer visible el avance del equipo y detectar problemas antes de la entrega final. El prototipado evolutivo complementa esta elección porque permite validar tempranamente el flujo de ventas, inventario y facturación con una versión funcional que será ampliada en cada sprint.",
        "Normal",
    ),
    (
        "PHP y Laravel fueron seleccionados por ofrecer una estructura basada en el patrón Modelo-Vista-Controlador, herramientas para validación, migraciones, relaciones de datos, seguridad y pruebas automatizadas. MySQL se utilizará por su compatibilidad con Laravel y XAMPP, además de proporcionar almacenamiento relacional adecuado para productos, ventas, facturas e inventario. Tailwind CSS permitirá construir una interfaz adaptable y consistente sin agregar una dependencia adicional al proyecto.",
        "Normal",
    ),
    (
        "Git y GitHub se utilizarán para el control de versiones, el respaldo del código y la identificación de los aportes individuales. El tablero de GitHub Projects permitirá organizar el Product Backlog y el Sprint Backlog, asignar responsables y dar seguimiento al estado de las tareas. Visual Studio Code, Composer, Node.js, Vite, XAMPP y phpMyAdmin complementarán el entorno de desarrollo e integración del sistema.",
        "Normal",
    ),
]

anchor = methodology_heading
for text, style in content:
    anchor = insert_paragraph_after(anchor, text, style)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
