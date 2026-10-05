from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\xampp\htdocs\Sistema-de-Facturaci-n-INS-\output\documents\Fase_1_metodologia_completada.docx")
OUTPUT = Path(r"C:\xampp\htdocs\Sistema-de-Facturaci-n-INS-\output\documents\Fase_1_metodologia_y_herramientas.docx")


def insert_after(anchor: Paragraph, text: str, style: str = "Normal", bullet: bool = False) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = "List Paragraph" if bullet else style
    if bullet:
        properties = paragraph._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number_id = OxmlElement("w:numId")
        number_id.set(qn("w:val"), "3")
        numbering.append(level)
        numbering.append(number_id)
        properties.append(numbering)
    paragraph.add_run(text)
    return paragraph


document = Document(SOURCE)

for paragraph in document.paragraphs:
    if "El tablero de GitHub Projects permitirá organizar" in paragraph.text:
        replacement = (
            "Jira se utilizará para organizar el Product Backlog y el Sprint Backlog, asignar responsables y dar seguimiento al estado de las tareas. "
            "Git y GitHub se utilizarán para el control de versiones, el respaldo del código y la identificación de los aportes individuales mediante ramas, commits y revisiones. "
            "Visual Studio Code, Composer, Node.js, Vite, XAMPP y phpMyAdmin complementarán el entorno de desarrollo e integración del sistema."
        )
        paragraph.clear()
        paragraph.add_run(replacement)

tools_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Herramientas de gestión.")

tools_content = [
    (
        "Para la planificación y el seguimiento del proyecto se utilizará Jira. Esta herramienta permitirá administrar el trabajo del equipo mediante un proyecto Scrum, manteniendo en un mismo espacio los requerimientos, historias de usuario, tareas, errores y actividades pendientes.",
        False,
    ),
    ("Configuración y uso de Jira", False),
    ("Product Backlog: contendrá las historias de usuario, requerimientos, mejoras y errores pendientes, ordenados según su prioridad.", True),
    ("Sprints: el trabajo priorizado se distribuirá en iteraciones de dos semanas, de acuerdo con la metodología definida para el proyecto.", True),
    ("Tablero Scrum: mostrará las tareas en los estados Pendiente, En proceso, En revisión y Finalizado.", True),
    ("Asignación de responsables: cada elemento tendrá un integrante responsable, prioridad, descripción y criterios de aceptación.", True),
    ("Seguimiento: los comentarios, evidencias y cambios de estado permitirán documentar el avance y los impedimentos encontrados.", True),
    ("Reportes: se consultará el avance del sprint para comparar el trabajo planificado con el trabajo completado.", True),
    (
        "Se seleccionó Jira porque ofrece soporte directo para proyectos Scrum, permite administrar el backlog y los sprints, y proporciona una visualización clara del flujo de trabajo. Su utilización favorecerá la coordinación del equipo, la transparencia de las responsabilidades y la evidencia de participación individual requerida para la evaluación del proyecto.",
        False,
    ),
    ("Control de versiones con Git y GitHub", False),
    (
        "Git se empleará como sistema de control de versiones y GitHub como repositorio remoto del código fuente. Cada integrante realizará sus cambios en ramas de trabajo identificadas con la tarea correspondiente de Jira. Los commits deberán describir de forma clara el cambio realizado y, antes de integrar una funcionalidad a la rama principal, se revisará que el código funcione y no afecte los módulos existentes.",
        False,
    ),
    ("El repositorio permitirá conservar el historial de cambios, respaldar el proyecto y verificar las contribuciones de cada integrante.", True),
    ("Las ramas separarán el desarrollo de nuevas funcionalidades y correcciones del código estable.", True),
    ("Los commits registrarán avances pequeños y relacionados con una tarea específica de Jira.", True),
    ("Las revisiones antes de integrar cambios ayudarán a mantener la calidad y reducir conflictos en el código.", True),
    ("La rama principal contendrá únicamente versiones revisadas y funcionales del sistema.", True),
    ("Relación entre Jira y GitHub", False),
    (
        "Cada historia o tarea creada en Jira tendrá una clave única. Esta clave se incluirá en el nombre de la rama y en los mensajes de commit cuando corresponda, permitiendo relacionar el avance técnico del repositorio con la actividad planificada. De esta forma, Jira mostrará qué debe realizarse y quién es responsable, mientras que GitHub conservará la evidencia del código desarrollado.",
        False,
    ),
]

anchor = tools_heading
for text, is_bullet in tools_content:
    style = "Heading 2" if text in {"Configuración y uso de Jira", "Control de versiones con Git y GitHub", "Relación entre Jira y GitHub"} else "Normal"
    anchor = insert_after(anchor, text, style=style, bullet=is_bullet)

sources_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Fuentes de información.")
source_anchor = sources_heading
for paragraph in document.paragraphs:
    if paragraph.text.strip().startswith("https://sv.corporacionisc.com"):
        source_anchor = paragraph
        break

source_anchor = insert_after(source_anchor, "Atlassian. (2026). Uso del backlog de Scrum en Jira Cloud.")
source_anchor = insert_after(source_anchor, "https://support.atlassian.com/jira-software-cloud/docs/use-your-scrum-backlog/")
source_anchor = insert_after(source_anchor, "GitHub Docs. (2026). Acerca de los repositorios.")
insert_after(source_anchor, "https://docs.github.com/en/repositories/creating-and-managing-repositories/about-repositories")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
